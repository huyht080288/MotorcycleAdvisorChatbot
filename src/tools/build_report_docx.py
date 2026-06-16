"""Build NHOM12_PHATTRIENHETHONGTHONGMINH_01.docx: giữ trang bìa, chèn nội dung + sơ đồ Mermaid dạng ảnh."""

from __future__ import annotations

import re
import subprocess
import sys
import tempfile
from pathlib import Path

REPORT_DIR = Path(__file__).resolve().parents[2] / "Report"
SOURCE_MD = REPORT_DIR / "NHOM12_PHATTRIENHETHONGTHONGMINH.md"
TEMPLATE_DOCX = REPORT_DIR / "NHOM12_PHATTRIENHETHONGTHONGMINH_01.docx"
OUTPUT_DOCX = TEMPLATE_DOCX
OUTPUT_FALLBACK = REPORT_DIR / "NHOM12_PHATTRIENHETHONGTHONGMINH_01_filled.docx"
DIAGRAM_DIR = REPORT_DIR / "images" / "diagrams"
PROCESSED_MD = REPORT_DIR / "_report_body.md"
CONTENT_DOCX = REPORT_DIR / "_report_body.docx"

# Bắt đầu nội dung từ MỤC LỤC (trang 2), bỏ phần bìa trùng trong .md
CONTENT_START_MARKER = "## MỤC LỤC"


def extract_body_markdown(source: str) -> str:
    idx = source.find(CONTENT_START_MARKER)
    if idx == -1:
        raise ValueError(f"Không tìm thấy '{CONTENT_START_MARKER}' trong file nguồn.")
    return source[idx:].strip() + "\n"


def render_mermaid_diagrams(markdown: str) -> str:
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    pattern = re.compile(r"```mermaid\n(.*?)```", re.DOTALL)
    blocks = pattern.findall(markdown)
    if not blocks:
        return markdown

    result = markdown
    for i, block in enumerate(blocks, start=1):
        name = f"diagram-{i:02d}"
        mmd_path = DIAGRAM_DIR / f"{name}.mmd"
        png_path = DIAGRAM_DIR / f"{name}.png"
        mmd_path.write_text(block.strip() + "\n", encoding="utf-8")

        cmd = [
            "npx",
            "-y",
            "@mermaid-js/mermaid-cli",
            "-i",
            str(mmd_path),
            "-o",
            str(png_path),
            "-b",
            "white",
            "-w",
            "1200",
        ]
        print(f"Rendering {name}...")
        subprocess.run(cmd, check=True, cwd=str(REPORT_DIR), shell=True)

        rel = f"images/diagrams/{name}.png"
        replacement = f"![Sơ đồ {i}]({rel})"
        result = result.replace(f"```mermaid\n{block}```", replacement, 1)

    return result


def pandoc_to_docx(md_path: Path, docx_path: Path) -> None:
    cmd = [
        "pandoc",
        str(md_path),
        "-o",
        str(docx_path),
        "--resource-path",
        str(REPORT_DIR),
        "-f",
        "markdown",
        "-t",
        "docx",
    ]
    subprocess.run(cmd, check=True)


def merge_docx(template: Path, content: Path, output: Path, fallback: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docxcompose.composer import Composer

    master = Document(str(template))
    if master.paragraphs:
        master.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    composer = Composer(master)
    composer.append(Document(str(content)))

    for target in (output, fallback):
        try:
            composer.save(str(target))
            return target
        except PermissionError:
            continue
    raise PermissionError(
        f"Không ghi được {output}. Đóng file Word đang mở và chạy lại script."
    )


def main() -> None:
    source = SOURCE_MD.read_text(encoding="utf-8")
    body = extract_body_markdown(source)
    body = render_mermaid_diagrams(body)
    PROCESSED_MD.write_text(body, encoding="utf-8")
    print(f"Wrote {PROCESSED_MD}")

    pandoc_to_docx(PROCESSED_MD, CONTENT_DOCX)
    print(f"Wrote {CONTENT_DOCX}")

    saved = merge_docx(TEMPLATE_DOCX, CONTENT_DOCX, OUTPUT_DOCX, OUTPUT_FALLBACK)
    print(f"Saved {saved}")


if __name__ == "__main__":
    try:
        main()
    except subprocess.CalledProcessError as exc:
        print(f"Command failed: {exc}", file=sys.stderr)
        sys.exit(exc.returncode)
